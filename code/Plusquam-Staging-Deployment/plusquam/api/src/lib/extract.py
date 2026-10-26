"""Safely extract plain text from an uploaded PDF or DOCX document.

Used by POST /projects to turn an optional upload into `sourceText` for the
assessment pipeline. Unsupported, mismatched, oversized, malformed, encrypted,
or unparseable documents raise ValueError so they cannot silently enter the
pipeline as trusted source material. Parser libraries are imported lazily so a
missing dependency for one format does not break the other.
"""

from __future__ import annotations

import io
import zipfile

MAX_UPLOAD_BYTES = 4 * 1024 * 1024
MAX_DOCX_ENTRIES = 1_000
MAX_DOCX_EXPANDED_BYTES = 20 * 1024 * 1024
MAX_PDF_PAGES = 100


def extract_text(filename: str, content_type: str, data: bytes) -> str:
    name = (filename or "").lower()
    ctype = (content_type or "").lower()

    if len(data) > MAX_UPLOAD_BYTES:
        raise ValueError("upload exceeds size limit")
    if name.endswith(".docx") and data.startswith(b"PK\x03\x04") and (
        not ctype or ctype in ("application/octet-stream", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    ):
        return _from_docx(data)
    if name.endswith(".pdf") and data.startswith(b"%PDF-") and ctype in ("", "application/pdf", "application/octet-stream"):
        return _from_pdf(data)
    raise ValueError("unsupported or invalid document")


def _from_docx(data: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_DOCX_ENTRIES:
                raise ValueError("too many DOCX archive entries")
            if "[Content_Types].xml" not in archive.namelist() or "word/document.xml" not in archive.namelist():
                raise ValueError("invalid DOCX structure")
            expanded = 0
            for info in infos:
                expanded += info.file_size
                if expanded > MAX_DOCX_EXPANDED_BYTES:
                    raise ValueError("DOCX expands beyond size limit")
                if info.file_size > 1024 * 1024 and info.compress_size and info.file_size / info.compress_size > 100:
                    raise ValueError("DOCX compression ratio exceeds limit")
    except (zipfile.BadZipFile, OSError) as exc:
        raise ValueError("invalid DOCX archive") from exc
    try:
        from docx import Document  # python-docx
    except Exception as e:  # pragma: no cover
        raise ValueError("DOCX parser unavailable") from e
    try:
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
        # Include table cell text too — specs often live in tables.
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts).strip()
    except Exception as e:
        raise ValueError("could not parse DOCX") from e


def _from_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except Exception as e:  # pragma: no cover
        raise ValueError("PDF parser unavailable") from e
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ValueError("encrypted PDF is unsupported")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise ValueError("PDF has too many pages")
        pages = [(page.extract_text() or "").strip() for page in reader.pages]
        return "\n\n".join(p for p in pages if p).strip()
    except Exception as e:
        raise ValueError("could not parse PDF") from e
